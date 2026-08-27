from pathlib import Path
import shutil
import openpyxl
from database import SessionLocal
import models

ROOT = Path(__file__).resolve().parent
BOOK = ROOT / "railway_inventory_import_2026-08-25.xlsx"

def main():
    wb = openpyxl.load_workbook(BOOK, read_only=True, data_only=True)
    ws = wb.worksheets[0]
    headers = [c.value for c in ws[4]]
    rows = list(ws.iter_rows(min_row=5, max_row=123, values_only=True))
    session = SessionLocal()
    added = 0
    try:
        existing = {
            ((p.name or "").strip().casefold(), float(p.price_usd or 0), float(p.weight_grams or 0))
            for p in session.query(models.Product).all()
        }
        for values in rows:
            data = dict(zip(headers, values))
            name = data.get("نام محصول")
            price = data.get("قیمت خرید (دلار)")
            if not name or price in (None, ""):
                continue
            weight = data.get("وزن (گرم)")
            key = (str(name).strip().casefold(), float(price), float(weight or 0))
            if key in existing:
                continue
            photo_paths = []
            for col in ("عکس ۱", "عکس ۲", "عکس ۳"):
                value = data.get(col)
                if not value:
                    continue
                source = ROOT / "photos" / Path(str(value)).name
                if source.exists():
                    photo_paths.append(str(source.relative_to(ROOT)).replace("\\", "/"))
            session.add(models.Product(
                name=str(name).strip(),
                description=data.get("توضیحات دستی"),
                price_usd=float(price),
                size=data.get("سایز"),
                weight_grams=float(weight) if weight not in (None, "") else None,
                location=None,
                category=data.get("دسته‌بندی") or "سایر",
                original_photo_path="|".join(photo_paths) or None,
            ))
            existing.add(key)
            added += 1
        session.commit()
        print(f"added={added} total={session.query(models.Product).count()}")
    finally:
        session.close()

if __name__ == "__main__":
    main()
